"""
reference_lists.py
------------------
Populates the Reference_Lists sheet and registers Excel DefinedNames
so that DataValidation dropdowns in other sheets can reference them
by name (e.g. ``=lstRiskLevels``).
"""

from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

_LISTS = {
    "Modules": ["Payables", "Receivables", "Procurement", "GL", "Projects", "SCM", "HCM"],
    "RiskLevels": ["Low", "Medium", "High", "Critical"],
    "Statuses": ["Draft", "Active", "Inactive", "Approved", "Rejected"],
    "YesNo": ["Y", "N"],
    "ReviewStatus": ["Pending Review", "Approved", "Rejected"],
}


def populate_reference_lists(wb, config: dict) -> None:  # noqa: ARG001
    """Write all reference lists to the Reference_Lists sheet and
    register each list as a workbook-level DefinedName.
    """
    ref_ws = wb["Reference_Lists"]

    for col_idx, (list_name, values) in enumerate(_LISTS.items(), start=1):
        # Column header
        ref_ws.cell(row=1, column=col_idx).value = list_name

        # Values starting at row 2
        for row_idx, val in enumerate(values, start=2):
            ref_ws.cell(row=row_idx, column=col_idx).value = val

        # Register a DefinedName for the data range (row 2 onward)
        col_letter = get_column_letter(col_idx)
        start = f"${col_letter}$2"
        end = f"${col_letter}${len(values) + 1}"
        wb.defined_names.add(
            DefinedName(
                f"lst{list_name}",
                attr_text=f"Reference_Lists!{start}:{end}",
            )
        )
