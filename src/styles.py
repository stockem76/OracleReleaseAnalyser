"""
styles.py
---------
All OpenPyXL style objects used across the workbook.

Call ``init_styles(config)`` once at startup (in main.py) before any
sheet-building module uses ``style_header``.
"""

from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

# Module-level style objects — populated by init_styles()
HEADER_FILL: PatternFill = None
HEADER_FONT: Font = None
CENTER: Alignment = None
THIN_BORDER: Border = None
CRITICAL_FILL: PatternFill = None
HIGH_FILL: PatternFill = None
MEDIUM_FILL: PatternFill = None
LOW_FILL: PatternFill = None


def init_styles(config: dict) -> None:
    """Initialise all module-level style objects from *config*.

    Must be called once before any sheet-building function is invoked.
    """
    global HEADER_FILL, HEADER_FONT, CENTER, THIN_BORDER
    global CRITICAL_FILL, HIGH_FILL, MEDIUM_FILL, LOW_FILL

    c = config["colours"]

    HEADER_FILL = PatternFill(
        start_color=c["header_fill"],
        end_color=c["header_fill"],
        fill_type="solid",
    )

    HEADER_FONT = Font(color="FFFFFF", bold=True)

    CENTER = Alignment(horizontal="center", vertical="center")

    thin = Side(style="thin", color="D9D9D9")
    THIN_BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

    CRITICAL_FILL = PatternFill(
        start_color=c["critical"],
        end_color=c["critical"],
        fill_type="solid",
    )

    HIGH_FILL = PatternFill(
        start_color=c["high"],
        end_color=c["high"],
        fill_type="solid",
    )

    MEDIUM_FILL = PatternFill(
        start_color=c["medium"],
        end_color=c["medium"],
        fill_type="solid",
    )

    LOW_FILL = PatternFill(
        start_color=c["low"],
        end_color=c["low"],
        fill_type="solid",
    )


def style_header(cell) -> None:
    """Apply the standard dark-blue header style to *cell*."""
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
    cell.alignment = CENTER
    cell.border = THIN_BORDER
