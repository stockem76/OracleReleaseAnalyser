"""Tests for src/seed_data.py, src/ingestion.py, CLI args, and auto-fit."""

import csv
import tempfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

from src.config_loader import load_config
from src import styles
from src.sheets import build_workbook, _auto_width
from src.reference_lists import populate_reference_lists
from src.seed_data import seed_all
from src.ingestion import ingest_csv, ingest_xlsx, _normalise_row, RELEASE_CHANGE_COLUMNS


# ---------------------------------------------------------------------------
# Auto-fit
# ---------------------------------------------------------------------------

def test_auto_width_min():
    assert _auto_width("ID") >= 12


def test_auto_width_max():
    assert _auto_width("A" * 100) <= 50


def test_auto_width_override():
    assert _auto_width("Change_Description") == 50
    assert _auto_width("Description") == 50


def test_auto_width_typical():
    w = _auto_width("Release_ID")
    assert 12 <= w <= 50


# ---------------------------------------------------------------------------
# Seed data
# ---------------------------------------------------------------------------

@pytest.fixture(scope="module")
def seeded_wb():
    cfg = load_config()
    styles.init_styles(cfg)
    wb = build_workbook(cfg)
    populate_reference_lists(wb, cfg)
    seed_all(wb)
    return wb


def test_clients_seeded(seeded_wb):
    ws = seeded_wb["Clients"]
    assert ws["A2"].value is not None, "Clients sheet should have data in row 2"


def test_release_changes_seeded(seeded_wb):
    ws = seeded_wb["Release_Changes"]
    assert ws["A2"].value is not None


def test_impact_assessments_seeded(seeded_wb):
    ws = seeded_wb["Impact_Assessments"]
    assert ws["A2"].value == "ASMT_001"
    assert ws["B2"].value == "CLT_DEMO01"


def test_severity_map_table_exists(seeded_wb):
    ws = seeded_wb["Scoring_Rules"]
    # tblSeverityMap headers in cols J/K
    assert ws["J1"].value == "Severity_Text"
    assert ws["K1"].value == "Score"
    # First data row
    assert ws["J2"].value == "Critical"
    assert ws["K2"].value == 5


def test_scoring_rules_seeded(seeded_wb):
    ws = seeded_wb["Scoring_Rules"]
    assert ws["A2"].value == "SR_001"


def test_raid_log_seeded(seeded_wb):
    ws = seeded_wb["RAID_Log"]
    assert ws["A2"].value == "RAID_001"


def test_release_readiness_seeded(seeded_wb):
    ws = seeded_wb["Release_Readiness"]
    assert ws["A2"].value is not None


# ---------------------------------------------------------------------------
# Ingestion — CSV
# ---------------------------------------------------------------------------

def test_ingest_csv_basic():
    with tempfile.NamedTemporaryFile(
        mode="w", suffix=".csv", delete=False, newline="", encoding="utf-8"
    ) as f:
        writer = csv.DictWriter(f, fieldnames=["Change_Title", "Oracle_Module", "Oracle_Severity"])
        writer.writeheader()
        writer.writerow({"Change_Title": "Test Feature", "Oracle_Module": "GL", "Oracle_Severity": "High"})
        tmp_path = f.name

    rows = ingest_csv(tmp_path, release_id="REL_TEST")
    assert len(rows) == 1
    assert rows[0]["Change_Title"] == "Test Feature"
    assert rows[0]["Oracle_Module"] == "GL"
    assert rows[0]["Release_ID"] == "REL_TEST"
    assert rows[0]["Change_ID"].startswith("CHG_")

    Path(tmp_path).unlink()


def test_ingest_csv_all_columns_present():
    with tempfile.NamedTemporaryFile(
        mode="w", suffix=".csv", delete=False, newline="", encoding="utf-8"
    ) as f:
        writer = csv.DictWriter(f, fieldnames=["title"])
        writer.writeheader()
        writer.writerow({"title": "Some change"})
        tmp_path = f.name

    rows = ingest_csv(tmp_path, release_id="REL_X")
    assert set(RELEASE_CHANGE_COLUMNS).issubset(set(rows[0].keys()))

    Path(tmp_path).unlink()


# ---------------------------------------------------------------------------
# Ingestion — XLSX
# ---------------------------------------------------------------------------

def test_ingest_xlsx_basic():
    from openpyxl import Workbook as WB

    tmp = Path(tempfile.mktemp(suffix=".xlsx"))
    wb = WB()
    ws = wb.active
    ws.append(["Change_Title", "Oracle_Module", "Oracle_Severity"])
    ws.append(["XLSX Feature", "HCM", "Medium"])
    wb.save(tmp)

    rows = ingest_xlsx(str(tmp), release_id="REL_XLSX_TEST")
    assert len(rows) == 1
    assert rows[0]["Change_Title"] == "XLSX Feature"
    assert rows[0]["Oracle_Module"] == "HCM"

    tmp.unlink()


# ---------------------------------------------------------------------------
# Normalise row — key casing tolerance
# ---------------------------------------------------------------------------

def test_normalise_row_case_insensitive():
    raw = {"CHANGE_TITLE": "Foo", "oracle_module": "ERP", "SEVERITY": "High"}
    row = _normalise_row(raw, "REL_001")
    assert row["Change_Title"] == "Foo"
    assert row["Oracle_Module"] == "ERP"
    assert row["Oracle_Severity"] == "High"


# ---------------------------------------------------------------------------
# CLI argument parser
# ---------------------------------------------------------------------------

def test_cli_parser_defaults():
    from src.main import _build_parser
    parser = _build_parser()
    args = parser.parse_args([])
    assert args.config is None
    assert args.output is None
    assert args.no_seed is False
    assert args.release is None
    assert args.client is None
    assert args.ingest_csv is None
    assert args.ingest_xlsx is None
    assert args.ingest_url is None


def test_cli_parser_flags():
    from src.main import _build_parser
    parser = _build_parser()
    args = parser.parse_args([
        "--config", "myconfig.yaml",
        "--output", "out/test.xlsx",
        "--no-seed",
        "--release", "REL_24C",
        "--client", "CLT_001",
    ])
    assert args.config == "myconfig.yaml"
    assert args.output == "out/test.xlsx"
    assert args.no_seed is True
    assert args.release == "REL_24C"
    assert args.client == "CLT_001"


def test_generate_no_seed(tmp_path):
    """generate(seed=False) should still produce a valid workbook."""
    import yaml
    from src.config_loader import _PROJECT_ROOT
    from src.main import generate

    base = _PROJECT_ROOT / "config" / "settings.yaml"
    with base.open() as fh:
        cfg = yaml.safe_load(fh)

    out = tmp_path / "no_seed.xlsx"
    cfg["output_file"] = str(out)
    cfg_file = tmp_path / "settings.yaml"
    with cfg_file.open("w") as fh:
        yaml.dump(cfg, fh)

    path = generate(str(cfg_file), seed=False)
    assert path.exists()
    wb = load_workbook(path)
    assert "Impact_Assessments" in wb.sheetnames
    # Row 2 col A should be empty (no seed)
    assert wb["Clients"]["A2"].value is None
