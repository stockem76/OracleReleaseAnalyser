"""Integration test — generates the full workbook and validates the result."""

import pytest
from pathlib import Path
from openpyxl import load_workbook

from src.main import generate
from src.sheets import SHEET_ORDER


@pytest.fixture(scope="module")
def generated_workbook(tmp_path_factory):
    """Generate the workbook once for the whole module into a temp dir."""
    import yaml
    from src.config_loader import _PROJECT_ROOT

    # Write a temp config that redirects output into pytest's tmp dir
    tmp_dir = tmp_path_factory.mktemp("output")
    out_file = tmp_dir / "test_output.xlsx"

    base_cfg_path = _PROJECT_ROOT / "config" / "settings.yaml"
    with base_cfg_path.open() as fh:
        cfg = yaml.safe_load(fh)

    cfg["output_file"] = str(out_file)

    tmp_cfg = tmp_dir / "settings.yaml"
    with tmp_cfg.open("w") as fh:
        yaml.dump(cfg, fh)

    output_path = generate(str(tmp_cfg))
    wb = load_workbook(output_path)
    return wb, output_path


def test_output_file_exists(generated_workbook):
    _, path = generated_workbook
    assert path.exists()


def test_sheet_count(generated_workbook):
    wb, _ = generated_workbook
    assert len(wb.sheetnames) >= 40


def test_all_expected_sheets_present(generated_workbook):
    wb, _ = generated_workbook
    for name in SHEET_ORDER:
        assert name in wb.sheetnames, f"Missing sheet: {name}"


def test_impact_assessments_has_headers(generated_workbook):
    wb, _ = generated_workbook
    ws = wb["Impact_Assessments"]
    assert ws["A1"].value == "Assessment_ID"


def test_reference_lists_populated(generated_workbook):
    wb, _ = generated_workbook
    ws = wb["Reference_Lists"]
    # First column header should be "Modules"
    assert ws["A1"].value == "Modules"
    # At least one value written below it
    assert ws["A2"].value is not None


def test_executive_summary_title(generated_workbook):
    wb, _ = generated_workbook
    ws = wb["Executive_Summary"]
    assert ws["A1"].value == "Oracle Fusion Impact Analyzer Dashboard"


def test_metadata_config_stamped(generated_workbook):
    wb, _ = generated_workbook
    ws = wb["Metadata_Config"]
    assert ws["A1"].value == "Generated_Date"
    assert ws["B1"].value is not None
